"""
reader.py — reads PDF, DOCX, and PPTX draft files and returns plain text.
"""

import os
from pathlib import Path


def read_pdf(path: str) -> str:
    import pymupdf
    doc = pymupdf.open(path)
    pages = []
    for page in doc:
        text = page.get_text()
        if text.strip():
            pages.append(text)
    doc.close()
    return "\n\n".join(pages)


def read_docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    # Also extract table text
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return "\n\n".join(paragraphs)


def read_pptx(path: str) -> str:
    from pptx import Presentation
    prs = Presentation(path)
    slides_text = []
    for i, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        parts.append(text)
        if parts:
            slides_text.append(f"[Slide {i}]\n" + "\n".join(parts))
    return "\n\n".join(slides_text)


def read_txt(path: str) -> str:
    encodings = ['utf-8', 'utf-8-sig', 'cp1256', 'latin-1']
    for enc in encodings:
        try:
            with open(path, 'r', encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, LookupError):
            continue
    with open(path, 'rb') as f:
        return f.read().decode('utf-8', errors='replace')


def read_file(path) -> str:
    """Auto-detect format and return extracted text."""
    path = str(path)
    ext = Path(path).suffix.lower()

    if ext == '.pdf':
        return read_pdf(path)
    elif ext in ('.docx', '.doc'):
        return read_docx(path)
    elif ext in ('.pptx', '.ppt'):
        return read_pptx(path)
    elif ext == '.txt':
        return read_txt(path)
    else:
        # Try as text first
        try:
            return read_txt(path)
        except Exception:
            raise ValueError(f"Unsupported file format: {ext}")
